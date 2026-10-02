"""Update FLAG IDs database from DLMS UA."""

from __future__ import annotations

import argparse
from difflib import unified_diff
from io import BytesIO
import json
from pathlib import Path
import sys
from typing import Final, cast
import urllib.error
import urllib.request

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

URL: Final = "https://www.dlms.com/srv/lib/Export_Flagids.php"
FILENAME: Final = "dlms_flagids.json"
COL: Final[dict[str, int]] = {
    "flag_id": 1,
    "manufacturer": 2,
    "country": 3,
    "region": 4,
}
OVERRIDES: Final[dict[str, str]] = {
    "KFM": "Shenzhen Kaifa Technology Co., Ltd.",
}

type FlagIdMap = dict[str, str]


def resolve_output_path(output_arg: Path | None) -> Path:
    """Resolve path to dlms_flagids.json."""
    if output_arg is not None:
        return output_arg

    cwd = Path.cwd()
    candidate = cwd / FILENAME
    if candidate.is_file():
        return candidate

    candidate = cwd / "custom_components" / "dlms_cosem" / FILENAME
    if candidate.is_file():
        return candidate

    script_dir = Path(__file__).resolve().parent
    repo_root = script_dir.parent
    candidate = repo_root / "custom_components" / "dlms_cosem" / FILENAME
    if candidate.is_file():
        return candidate

    return Path("custom_components") / "dlms_cosem" / FILENAME


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    """Parse command line arguments."""
    parser = argparse.ArgumentParser(
        description="Update FLAG IDs database from DLMS UA."
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="Path to dlms_flagids.json output file (default: auto-detected)",
    )
    parser.add_argument(
        "-u",
        "--url",
        type=str,
        default=URL,
        help=f"URL to export FLAG IDs spreadsheet (default: {URL})",
    )
    parser.add_argument(
        "-c",
        "--check",
        action="store_true",
        help="Check if database is up to date without modifying files",
    )
    parser.add_argument(
        "--diff",
        action="store_true",
        help="Show unified diff of changes",
    )
    return parser.parse_args(argv)


def fetch_spreadsheet(url: str) -> bytes:
    """Download the spreadsheet from DLMS UA."""
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (compatible; DLMS-FLAG-Sync/1.0; "
            "+https://github.com/denpamusic/homeassistant-dlms-cosem)"
        )
    }
    request = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(request, timeout=30) as response:
        return cast(bytes, response.read())


def parse_worksheet(ws: Worksheet) -> FlagIdMap:
    """Parse flag id and manufacturer pairs from worksheet."""
    rows: list[tuple[str, str]] = []
    for row in range(2, ws.max_row + 1):
        flag_cell = ws.cell(row, COL["flag_id"]).value
        mfg_cell = ws.cell(row, COL["manufacturer"]).value
        if flag_cell and mfg_cell:
            flag_id = str(flag_cell).strip()
            manufacturer = str(mfg_cell).strip()
            if flag_id and manufacturer:
                rows.append((flag_id, manufacturer))

    # Sort rows so duplicate FLAG IDs resolve deterministically regardless of server order
    rows.sort(key=lambda item: (item[0], item[1]))
    return dict(rows)


def apply_overrides(
    manufacturers: FlagIdMap, overrides: dict[str, str]
) -> list[tuple[str, str, str]]:
    """Apply manufacturer overrides and return replacement details."""
    replacements: list[tuple[str, str, str]] = []
    for flag_id, override in overrides.items():
        if (current := manufacturers.get(flag_id)) and current != override:
            manufacturers[flag_id] = override
            replacements.append((flag_id, current, override))
    return replacements


def compare_flag_ids(
    old_data: FlagIdMap, new_data: FlagIdMap
) -> tuple[list[str], list[str], list[tuple[str, str, str]]]:
    """Compare old and new flag id dictionaries."""
    old_keys = set(old_data.keys())
    new_keys = set(new_data.keys())

    added = sorted(new_keys - old_keys)
    removed = sorted(old_keys - new_keys)
    changed = [
        (k, old_data[k], new_data[k])
        for k in sorted(old_keys & new_keys)
        if old_data[k] != new_data[k]
    ]

    return added, removed, changed


def report_changes(
    output_path: Path,
    added: list[str],
    removed: list[str],
    changed: list[tuple[str, str, str]],
    new_data: FlagIdMap,
    check_only: bool,
) -> None:
    """Print update details."""
    sys.stdout.write(f"{output_path.name}:\n")
    if not added and not removed and not changed:
        sys.stdout.write("  In sync (no changes).\n\n")
        return

    action_add = "Would add" if check_only else "Added"
    action_rem = "Would remove" if check_only else "Removed"
    action_chg = "Would update" if check_only else "Updated"

    if added:
        sys.stdout.write(f"  {action_add} ({len(added)} keys):\n")
        for key in added:
            sys.stdout.write(f"    + {key}: {new_data[key]}\n")
    if removed:
        sys.stdout.write(f"  {action_rem} ({len(removed)} keys):\n")
        for key in removed:
            sys.stdout.write(f"    - {key}\n")
    if changed:
        sys.stdout.write(f"  {action_chg} ({len(changed)} keys):\n")
        for key, old_val, new_val in changed:
            sys.stdout.write(f'    ~ {key}: "{old_val}" -> "{new_val}"\n')
    sys.stdout.write("\n")


def main(argv: list[str] | None = None) -> int:
    """Run FLAG IDs database update tool."""
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8")

    args = parse_args(argv)
    output_path = resolve_output_path(args.output)

    old_data: FlagIdMap = {}
    if output_path.is_file():
        try:
            with output_path.open("r", encoding="utf-8") as f:
                old_data = json.load(f)
        except json.JSONDecodeError as err:
            sys.stderr.write(f"Error: Failed to parse '{output_path}': {err}\n")
            return 1

    mode_prefix = "[CHECK] " if args.check else ""
    sys.stdout.write(f"{mode_prefix}Updating FLAG IDs from '{args.url}'...\n\n")

    try:
        spreadsheet_bytes = fetch_spreadsheet(args.url)
    except urllib.error.URLError as err:
        sys.stderr.write(f"Error: Failed to fetch spreadsheet: {err}\n")
        return 1

    try:
        workbook = openpyxl.load_workbook(BytesIO(spreadsheet_bytes))
        manufacturers = parse_worksheet(workbook.active)
        workbook.close()
    except Exception as err:
        sys.stderr.write(f"Error: Failed to parse spreadsheet: {err}\n")
        return 1

    apply_overrides(manufacturers, OVERRIDES)
    sorted_manufacturers = dict(sorted(manufacturers.items()))

    added, removed, changed = compare_flag_ids(old_data, sorted_manufacturers)
    report_changes(
        output_path=output_path,
        added=added,
        removed=removed,
        changed=changed,
        new_data=sorted_manufacturers,
        check_only=args.check,
    )

    old_dump = json.dumps(old_data, indent=2, ensure_ascii=False) + "\n"
    new_dump = json.dumps(sorted_manufacturers, indent=2, ensure_ascii=False) + "\n"
    has_changes = old_dump != new_dump

    if args.diff and has_changes:
        sys.stdout.write("Diff:\n")
        sys.stdout.writelines(
            unified_diff(
                old_dump.splitlines(keepends=True),
                new_dump.splitlines(keepends=True),
                fromfile=f"{output_path.name}-old",
                tofile=f"{output_path.name}-new",
                n=1,
            )
        )
        sys.stdout.write("\n")

    summary_title = "[CHECK] Summary:" if args.check else "Summary:"
    sys.stdout.write(f"{summary_title}\n")
    if not has_changes:
        sys.stdout.write(
            f"  {output_path.name}: in sync ({len(sorted_manufacturers)} entries)\n"
        )
    else:
        status_suffix = "to update" if args.check else "updated"
        sys.stdout.write(
            f"  {output_path.name}: +{len(added)}, -{len(removed)}, ~{len(changed)} "
            f"({status_suffix}, total {len(sorted_manufacturers)} entries)\n"
        )

    if not args.check and has_changes:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("w", encoding="utf-8", newline="\n") as f:
            f.write(new_dump)

    if args.check:
        if has_changes:
            sys.stdout.write("Check failed: FLAG IDs database is out of date.\n")
            return 1
        sys.stdout.write("Check passed: FLAG IDs database is up to date.\n")
        return 0

    return 0


if __name__ == "__main__":
    sys.exit(main())
